# -*- coding: utf-8 -*-
"""Markdown -> DOCX converter (tables, code blocks, headings, lists, quotes)."""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def strip_inline(s):
    """Remove markdown inline markers, return plain text."""
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    s = re.sub(r'`([^`]+)`', r'\1', s)
    s = re.sub(r'\*(.+?)\*', r'\1', s)
    s = re.sub(r'\[(.+?)\]\([^)]*\)', r'\1', s)
    return s


def add_rich(p, text, size=10.5, bold=False, mono=False):
    """Add text with **bold** segments rendered bold."""
    font = 'Consolas' if mono else 'Microsoft YaHei'
    for part in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        else:
            r = p.add_run(strip_inline(part))
            r.bold = bold
        r.font.size = Pt(size)
        r.font.name = font
        r._element.rPr.rFonts.set(qn('w:eastAsia'), font)


def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Microsoft YaHei'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # code fence
        if line.strip().startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run('\n'.join(buf))
            r.font.name = 'Consolas'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            continue

        # horizontal rule
        if re.fullmatch(r'-{3,}|\*{3,}|_{3,}', line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run('─' * 60)
            r.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            r.font.size = Pt(8)
            i += 1
            continue

        # table
        if line.strip().startswith('|') and i + 1 < n and re.match(r'^\s*\|[\s:\-|]+\|\s*$', lines[i + 1]):
            def cells(row):
                row = row.strip()
                if row.startswith('|'):
                    row = row[1:]
                if row.endswith('|'):
                    row = row[:-1]
                return [c.strip() for c in row.split('|')]
            header = cells(line)
            i += 2
            body = []
            while i < n and lines[i].strip().startswith('|'):
                body.append(cells(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = t.rows[0].cells
            for c, txt in enumerate(header):
                hdr[c].text = ''
                p = hdr[c].paragraphs[0]
                add_rich(p, txt, size=9.5, bold=True)
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shd = hdr[c]._tc.get_or_add_tcPr()
                from docx.oxml import OxmlElement
                el = OxmlElement('w:shd')
                el.set(qn('w:fill'), '4472C4')
                shd.append(el)
            for row in body:
                rowcells = t.add_row().cells
                for c in range(len(header)):
                    txt = row[c] if c < len(row) else ''
                    rowcells[c].text = ''
                    add_rich(rowcells[c].paragraphs[0], txt, size=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # heading
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=min(lvl, 4))
            add_rich(h, m.group(2), size=max(16 - lvl * 1.5, 11), bold=True)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            i += 1
            continue

        # blockquote
        if line.strip().startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                buf.append(re.sub(r'^\s*>\s?', '', lines[i]))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_rich(p, ' '.join(x for x in buf if x.strip()), size=10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                r.italic = True
            continue

        # list
        m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        if m:
            indent = len(m.group(1)) // 2
            txt = m.group(3)
            p = doc.add_paragraph(style='List Bullet' if m.group(2) in '-*' else 'List Number')
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * indent)
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, txt)
            i += 1
            continue

        # blank
        if not line.strip():
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_rich(p, line)
        i += 1

    doc.save(docx_path)
    print('saved:', docx_path)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
